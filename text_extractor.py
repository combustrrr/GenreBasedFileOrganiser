"""
Text extraction module for various Office file formats.
Supports .docx, .xlsx, and .pptx files.
"""

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
import os


class TextExtractor:
    """Extract text content from various Office file formats."""
    
    def __init__(self):
        """Initialize the TextExtractor."""
        self.supported_formats = {'.docx', '.xlsx', '.pptx'}
    
    def is_supported(self, file_path):
        """Check if file format is supported.
        
        Args:
            file_path: Path to the file
            
        Returns:
            bool: True if file format is supported
        """
        _, ext = os.path.splitext(file_path)
        return ext.lower() in self.supported_formats
    
    def extract_from_docx(self, file_path):
        """Extract text from Word (.docx) file.
        
        Args:
            file_path: Path to the .docx file
            
        Returns:
            str: Extracted text content
        """
        try:
            doc = Document(file_path)
            text_parts = []
            
            # Extract text from paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
            
            return ' '.join(text_parts)
        except Exception as e:
            print(f"Error extracting text from {file_path}: {e}")
            return ""
    
    def extract_from_xlsx(self, file_path):
        """Extract text from Excel (.xlsx) file.
        
        Args:
            file_path: Path to the .xlsx file
            
        Returns:
            str: Extracted text content
        """
        try:
            wb = load_workbook(file_path, data_only=True)
            text_parts = []
            
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None and str(cell).strip():
                            text_parts.append(str(cell))
            
            return ' '.join(text_parts)
        except Exception as e:
            print(f"Error extracting text from {file_path}: {e}")
            return ""
    
    def extract_from_pptx(self, file_path):
        """Extract text from PowerPoint (.pptx) file.
        
        Args:
            file_path: Path to the .pptx file
            
        Returns:
            str: Extracted text content
        """
        try:
            prs = Presentation(file_path)
            text_parts = []
            
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        text_parts.append(shape.text)
            
            return ' '.join(text_parts)
        except Exception as e:
            print(f"Error extracting text from {file_path}: {e}")
            return ""
    
    def extract_text(self, file_path):
        """Extract text from supported file formats.
        
        Args:
            file_path: Path to the file
            
        Returns:
            str: Extracted text content
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()
        
        if ext == '.docx':
            return self.extract_from_docx(file_path)
        elif ext == '.xlsx':
            return self.extract_from_xlsx(file_path)
        elif ext == '.pptx':
            return self.extract_from_pptx(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")
