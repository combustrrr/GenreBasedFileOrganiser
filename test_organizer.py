"""
Basic tests for the Genre-Based File Organizer.
Note: These are minimal tests to verify core functionality.
"""

import unittest
import os
import tempfile
import shutil
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

# Import our modules
from text_extractor import TextExtractor
from embedding_generator import EmbeddingGenerator
from clusterer import DocumentClusterer


class TestTextExtractor(unittest.TestCase):
    """Test the text extraction functionality."""
    
    def setUp(self):
        """Set up test fixtures."""
        self.extractor = TextExtractor()
        self.test_dir = tempfile.mkdtemp()
    
    def tearDown(self):
        """Clean up test files."""
        shutil.rmtree(self.test_dir, ignore_errors=True)
    
    def test_supported_formats(self):
        """Test that correct formats are supported."""
        self.assertTrue(self.extractor.is_supported("test.docx"))
        self.assertTrue(self.extractor.is_supported("test.xlsx"))
        self.assertTrue(self.extractor.is_supported("test.pptx"))
        self.assertFalse(self.extractor.is_supported("test.txt"))
        self.assertFalse(self.extractor.is_supported("test.pdf"))
    
    def test_extract_from_docx(self):
        """Test extracting text from Word document."""
        # Create a test docx file
        doc = Document()
        doc.add_paragraph("This is a test document.")
        doc.add_paragraph("It contains multiple paragraphs.")
        
        test_file = os.path.join(self.test_dir, "test.docx")
        doc.save(test_file)
        
        # Extract text
        text = self.extractor.extract_from_docx(test_file)
        
        self.assertIn("test document", text)
        self.assertIn("multiple paragraphs", text)
    
    def test_extract_from_xlsx(self):
        """Test extracting text from Excel file."""
        # Create a test xlsx file
        wb = Workbook()
        ws = wb.active
        ws['A1'] = "Product"
        ws['B1'] = "Price"
        ws['A2'] = "Apple"
        ws['B2'] = 1.50
        
        test_file = os.path.join(self.test_dir, "test.xlsx")
        wb.save(test_file)
        
        # Extract text
        text = self.extractor.extract_from_xlsx(test_file)
        
        self.assertIn("Product", text)
        self.assertIn("Apple", text)
    
    def test_extract_from_pptx(self):
        """Test extracting text from PowerPoint file."""
        # Create a test pptx file
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Test Presentation"
        
        test_file = os.path.join(self.test_dir, "test.pptx")
        prs.save(test_file)
        
        # Extract text
        text = self.extractor.extract_from_pptx(test_file)
        
        self.assertIn("Test Presentation", text)


class TestEmbeddingGenerator(unittest.TestCase):
    """Test the embedding generation functionality."""
    
    @classmethod
    def setUpClass(cls):
        """Set up class-level fixtures (load model once)."""
        print("\nLoading DistilBERT model for testing...")
        cls.generator = EmbeddingGenerator()
    
    def test_generate_embedding(self):
        """Test generating embedding for text."""
        text = "This is a test document about machine learning."
        embedding = self.generator.generate_embedding(text)
        
        # Check embedding shape (DistilBERT outputs 768-dim vectors)
        self.assertEqual(embedding.shape, (768,))
        
        # Check that embedding is not all zeros
        self.assertNotEqual(embedding.sum(), 0)
    
    def test_empty_text_embedding(self):
        """Test generating embedding for empty text."""
        embedding = self.generator.generate_embedding("")
        
        # Should return zero vector for empty text
        self.assertEqual(embedding.shape, (768,))
        self.assertEqual(embedding.sum(), 0)
    
    def test_similar_texts_similar_embeddings(self):
        """Test that similar texts produce similar embeddings."""
        text1 = "Machine learning is a subset of artificial intelligence."
        text2 = "AI and machine learning are closely related fields."
        text3 = "I love eating pizza and pasta for dinner."
        
        emb1 = self.generator.generate_embedding(text1)
        emb2 = self.generator.generate_embedding(text2)
        emb3 = self.generator.generate_embedding(text3)
        
        # Compute cosine similarities
        import numpy as np
        
        def cosine_similarity(a, b):
            return np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b))
        
        sim_12 = cosine_similarity(emb1, emb2)
        sim_13 = cosine_similarity(emb1, emb3)
        
        # Similar texts should be more similar than dissimilar texts
        self.assertGreater(sim_12, sim_13)


class TestDocumentClusterer(unittest.TestCase):
    """Test the clustering functionality."""
    
    def test_cluster_documents(self):
        """Test clustering documents."""
        import numpy as np
        
        # Create some synthetic embeddings
        # Group 1: around [1, 0, 0, ...]
        # Group 2: around [0, 1, 0, ...]
        embeddings = []
        for i in range(5):
            emb = np.zeros(768)
            emb[0] = 1.0 + np.random.normal(0, 0.1)
            embeddings.append(emb)
        for i in range(5):
            emb = np.zeros(768)
            emb[1] = 1.0 + np.random.normal(0, 0.1)
            embeddings.append(emb)
        
        embeddings = np.array(embeddings)
        
        clusterer = DocumentClusterer()
        clusters = clusterer.cluster_documents(embeddings, n_clusters=2)
        
        # Should have 2 clusters
        unique_clusters = set(clusters)
        self.assertEqual(len(unique_clusters), 2)
        
        # First 5 should be in one cluster, last 5 in another
        # (or vice versa, doesn't matter which label)
        cluster_0 = clusters[0]
        for i in range(5):
            self.assertEqual(clusters[i], cluster_0)
    
    def test_build_index(self):
        """Test building FAISS index."""
        import numpy as np
        
        embeddings = np.random.randn(10, 768).astype('float32')
        
        clusterer = DocumentClusterer()
        clusterer.build_index(embeddings)
        
        self.assertIsNotNone(clusterer.index)
        self.assertEqual(clusterer.index.ntotal, 10)


def run_tests():
    """Run all tests."""
    unittest.main(argv=[''], verbosity=2, exit=False)


if __name__ == "__main__":
    run_tests()
